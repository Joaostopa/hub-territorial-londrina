from io import BytesIO
import pandas as pd

def read_table(name,content):
    if len(content)>25_000_000:
        raise ValueError('Limite de 25 MB por arquivo.')
    if name.lower().endswith('.xlsx'):
        return pd.read_excel(BytesIO(content),dtype={'codigo':str,'id_externo':str})
    if name.lower().endswith('.json'):
        import json
        payload=json.loads(content)
        if not isinstance(payload,list):
            raise ValueError('O JSON normalizado deve ser uma lista de registros.')
        return pd.DataFrame(payload)
    return pd.read_csv(BytesIO(content),sep=None,engine='python',dtype={'codigo':str,'id_externo':str})

def safe_export(df):
    result=df.copy()
    for col in result.select_dtypes(include=['object','string']).columns:
        result[col]=result[col].map(lambda v: "'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@','\t','\r')) else v)
    return result

def csv_bytes(df):
    return safe_export(df).to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')

def excel_bytes(sheets):
    buffer=BytesIO()
    with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
        for name,frame in sheets.items():
            frame=safe_export(frame)
            frame.to_excel(writer,sheet_name=name[:31],index=False)
            ws=writer.sheets[name[:31]]
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.font=Font(color='FFFFFF',bold=True)
                cell.fill=PatternFill('solid',fgColor='123F50')
            for cells in ws.columns:
                letter=cells[0].column_letter
                ws.column_dimensions[letter].width=min(48,max(14,len(str(cells[0].value))+3))
    return buffer.getvalue()
